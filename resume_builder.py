#!/usr/bin/env python3
"""Build a truthful ATS-safe resume from a locked master fact bank.

The builder is intentionally deterministic and non-generative:
- job-description skills are ranked with resume_tailor.py;
- only bullets marked verified=true can be selected;
- bullet text is copied verbatim from the master file;
- employer/title/date/education facts must be locked=true and are never rewritten;
- the output DOCX uses one column, no tables/text boxes, no headers/footers,
  no hidden text, and no watermark-like drawing objects.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from pathlib import Path
from typing import Iterable

import requests
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from resume_tailor import (
    DEFAULT_INVENTORY,
    load_inventory,
    rank_jd_skills,
    select_unverified_gaps,
    select_verified_skills,
)


REQUEST_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/152.0 Safari/537.36"
    ),
    "Accept-Language": "en-US,en;q=0.9",
}

STOPWORDS = {
    "about", "after", "also", "and", "are", "been", "being", "but", "can",
    "company", "experience", "for", "from", "have", "into", "job", "more",
    "our", "role", "that", "the", "their", "this", "through", "using", "with",
    "will", "work", "working", "you", "your", "years", "preferred", "required",
    "requirements", "qualifications", "minimum", "skills", "team", "teams",
}

QUALIFICATION_MARKERS = (
    "required",
    "requirements",
    "minimum qualifications",
    "basic qualifications",
    "preferred qualifications",
    "must have",
    "experience with",
    "years of experience",
    "proficiency",
    "proficient",
    "degree",
    "bachelor",
    "master",
)


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def _normalize_key(text: str) -> str:
    return re.sub(r"[^a-z0-9+#.]+", " ", _clean(text).casefold()).strip()


def _phrase_present(text: str, phrase: str) -> bool:
    phrase = _clean(phrase)
    if not phrase:
        return False
    pattern = re.compile(
        rf"(?<![A-Za-z0-9]){re.escape(phrase)}(?![A-Za-z0-9])",
        flags=re.I,
    )
    return bool(pattern.search(text))


def load_master(path: str | Path) -> dict:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    validate_master(data)
    return data


def validate_master(master: dict) -> None:
    """Fail closed if mutable/unverified resume facts are supplied."""
    if not isinstance(master, dict):
        raise ValueError("master resume must be a JSON object")
    contact = master.get("contact")
    if not isinstance(contact, dict) or not _clean(contact.get("name", "")):
        raise ValueError("master.contact.name is required")

    ids: set[str] = set()
    for section in ("experience", "projects"):
        entries = master.get(section, [])
        if not isinstance(entries, list):
            raise ValueError(f"master.{section} must be a list")
        for entry in entries:
            if not isinstance(entry, dict):
                raise ValueError(f"every {section} entry must be an object")
            if entry.get("locked") is not True:
                raise ValueError(f"every {section} entry must set locked=true")
            entry_id = _clean(entry.get("id", ""))
            if not entry_id:
                raise ValueError(f"every {section} entry needs a stable id")
            if entry_id in ids:
                raise ValueError(f"duplicate resume id: {entry_id}")
            ids.add(entry_id)

            bullets = entry.get("bullets", [])
            if not isinstance(bullets, list):
                raise ValueError(f"{entry_id}.bullets must be a list")
            for bullet in bullets:
                if not isinstance(bullet, dict):
                    raise ValueError(f"{entry_id} has a non-object bullet")
                bullet_id = _clean(bullet.get("id", ""))
                if not bullet_id:
                    raise ValueError(f"every bullet under {entry_id} needs an id")
                if bullet_id in ids:
                    raise ValueError(f"duplicate resume id: {bullet_id}")
                ids.add(bullet_id)
                if not _clean(bullet.get("text", "")):
                    raise ValueError(f"{bullet_id} has no text")

    education = master.get("education", [])
    if not isinstance(education, list):
        raise ValueError("master.education must be a list")
    for item in education:
        if not isinstance(item, dict) or item.get("locked") is not True:
            raise ValueError("every education entry must be an object with locked=true")


def locked_fact_digest(master: dict) -> str:
    """Hash the immutable identity/date fields for auditability."""
    payload = {
        "contact": master.get("contact", {}),
        "summary": master.get("summary", ""),
        "experience": [
            {
                key: item.get(key, "")
                for key in ("id", "company", "title", "location", "start", "end", "locked")
            }
            for item in master.get("experience", [])
        ],
        "projects": [
            {
                key: item.get(key, "")
                for key in ("id", "name", "subtitle", "locked")
            }
            for item in master.get("projects", [])
        ],
        "education": master.get("education", []),
        "certifications": master.get("certifications", []),
    }
    raw = json.dumps(payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def fetch_jd_text(url: str, timeout: int = 20) -> str:
    """Fetch readable visible text from a public job page."""
    response = requests.get(url, headers=REQUEST_HEADERS, timeout=timeout)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    for node in soup(["script", "style", "noscript", "svg"]):
        node.decompose()
    text = "\n".join(_clean(line) for line in soup.stripped_strings)
    if len(text) < 120:
        raise ValueError("job page did not expose enough readable text; save the JD to a text file instead")
    return text


def _candidate_lines(jd_text: str) -> list[str]:
    lines: list[str] = []
    for raw in jd_text.splitlines():
        cleaned = _clean(re.sub(r"^[\s•*\-–—]+", "", raw))
        if cleaned:
            lines.append(cleaned)
    if len(lines) < 8:
        lines.extend(
            _clean(part)
            for part in re.split(r"(?<=[.!?])\s+", jd_text)
            if 20 <= len(_clean(part)) <= 420
        )
    return list(dict.fromkeys(lines))


def extract_top_qualification_lines(jd_text: str, limit: int = 12) -> list[dict]:
    """Return deterministic high-signal JD qualification lines for the report."""
    ranked: list[dict] = []
    for line in _candidate_lines(jd_text):
        lower = line.casefold()
        score = 0
        matched_markers = []
        for marker in QUALIFICATION_MARKERS:
            if marker in lower:
                score += 5 if marker not in ("preferred qualifications",) else 3
                matched_markers.append(marker)
        if re.search(r"\b\d+\+?\s*(?:years?|yrs?)\b", lower):
            score += 4
        if re.search(r"\b(?:python|sql|java|aws|azure|gcp|spark|machine learning|software|data)\b", lower):
            score += 2
        if 25 <= len(line) <= 300:
            score += 1
        if score:
            ranked.append({"text": line, "score": score, "markers": matched_markers})
    ranked.sort(key=lambda x: (-x["score"], len(x["text"]), x["text"].casefold()))
    return ranked[:limit]


def _jd_terms(jd_text: str) -> set[str]:
    tokens = re.findall(r"[A-Za-z][A-Za-z0-9+#.\-/]{2,}", jd_text.casefold())
    return {token for token in tokens if len(token) >= 4 and token not in STOPWORDS}


def _skill_score_map(ranked_skills: Iterable[dict]) -> dict[str, float]:
    return {_normalize_key(item["name"]): float(item.get("score", 0)) for item in ranked_skills}


def score_bullet(bullet: dict, jd_text: str, ranked_skills: list[dict]) -> dict:
    """Score one existing bullet without changing its wording."""
    if bullet.get("verified") is not True:
        return {
            "id": bullet.get("id", ""),
            "score": -1.0,
            "eligible": False,
            "matched_skills": [],
            "matched_keywords": [],
            "term_overlap": [],
        }

    skill_scores = _skill_score_map(ranked_skills)
    matched_skills: list[str] = []
    skill_score = 0.0
    for skill in bullet.get("skills", []):
        key = _normalize_key(skill)
        if key in skill_scores:
            matched_skills.append(str(skill))
            skill_score += skill_scores[key] * 4.0

    matched_keywords = [
        str(keyword)
        for keyword in bullet.get("keywords", [])
        if _phrase_present(jd_text, str(keyword))
    ]
    keyword_score = len(matched_keywords) * 3.0

    bullet_terms = _jd_terms(_clean(bullet.get("text", "")))
    overlap = sorted(bullet_terms & _jd_terms(jd_text))
    term_score = min(len(overlap), 8) * 0.5

    score = skill_score + keyword_score + term_score
    return {
        "id": bullet.get("id", ""),
        "score": round(score, 3),
        "eligible": True,
        "matched_skills": matched_skills,
        "matched_keywords": matched_keywords,
        "term_overlap": overlap[:12],
    }


def select_entry_bullets(
    entry: dict,
    jd_text: str,
    ranked_skills: list[dict],
    *,
    max_bullets: int,
    min_bullets: int = 1,
) -> tuple[list[dict], list[dict]]:
    scored: list[tuple[int, dict, dict]] = []
    for index, bullet in enumerate(entry.get("bullets", [])):
        detail = score_bullet(bullet, jd_text, ranked_skills)
        if detail["eligible"]:
            scored.append((index, bullet, detail))

    scored.sort(key=lambda item: (-item[2]["score"], item[0]))
    positives = [item for item in scored if item[2]["score"] > 0]
    chosen = positives[:max_bullets]
    if len(chosen) < min_bullets:
        chosen_ids = {item[1].get("id") for item in chosen}
        for item in scored:
            if item[1].get("id") not in chosen_ids:
                chosen.append(item)
                chosen_ids.add(item[1].get("id"))
            if len(chosen) >= min_bullets:
                break
    chosen = chosen[:max_bullets]
    return [item[1] for item in chosen], [item[2] for item in chosen]


def select_projects(
    projects: list[dict],
    jd_text: str,
    ranked_skills: list[dict],
    *,
    max_projects: int = 3,
    max_bullets_per_project: int = 3,
) -> list[dict]:
    candidates: list[dict] = []
    for index, project in enumerate(projects):
        bullets, details = select_entry_bullets(
            project,
            jd_text,
            ranked_skills,
            max_bullets=max_bullets_per_project,
            min_bullets=0,
        )
        score = sum(max(0.0, item["score"]) for item in details)
        name_blob = f"{project.get('name', '')} {project.get('subtitle', '')}"
        name_overlap = len(_jd_terms(name_blob) & _jd_terms(jd_text))
        score += name_overlap
        if score <= 0:
            continue
        candidates.append(
            {
                "index": index,
                "entry": project,
                "bullets": bullets,
                "bullet_scores": details,
                "score": round(score, 3),
            }
        )
    candidates.sort(key=lambda item: (-item["score"], item["index"]))
    return candidates[:max_projects]


def build_selection(
    master: dict,
    jd_text: str,
    inventory: list[dict],
    *,
    max_skills: int = 15,
    max_bullets_per_experience: int = 4,
    max_projects: int = 3,
    max_bullets_per_project: int = 3,
) -> dict:
    ranked_skills = rank_jd_skills(jd_text, inventory)
    selected_skills = select_verified_skills(ranked_skills, max_skills=max_skills)
    gaps = select_unverified_gaps(ranked_skills, max_items=max_skills)

    experience_selection = []
    for entry in master.get("experience", []):
        bullets, details = select_entry_bullets(
            entry,
            jd_text,
            ranked_skills,
            max_bullets=max_bullets_per_experience,
            min_bullets=1,
        )
        experience_selection.append(
            {
                "entry": entry,
                "bullets": bullets,
                "bullet_scores": details,
            }
        )

    project_selection = select_projects(
        master.get("projects", []),
        jd_text,
        ranked_skills,
        max_projects=max_projects,
        max_bullets_per_project=max_bullets_per_project,
    )

    return {
        "locked_fact_sha256": locked_fact_digest(master),
        "selected_verified_skills": selected_skills,
        "requested_but_unverified": gaps,
        "ranked_jd_skills": ranked_skills,
        "top_qualification_lines": extract_top_qualification_lines(jd_text),
        "experience": experience_selection,
        "projects": project_selection,
        "limits": {
            "max_skills": max_skills,
            "max_bullets_per_experience": max_bullets_per_experience,
            "max_projects": max_projects,
            "max_bullets_per_project": max_bullets_per_project,
        },
        "notes": [
            "All rendered experience/project bullets are copied verbatim from verified master bullets.",
            "Employer names, titles, dates, education and contact facts are never rewritten.",
            "Unverified JD skills are reported as gaps and are not inserted.",
            "No generative-AI API is called by the builder.",
        ],
    }


def _set_run_font(run, *, size: float = 10.0, bold: bool | None = None) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _add_section_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper())
    _set_run_font(run, size=10.5, bold=True)


def _add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(f"- {_clean(text)}")
    _set_run_font(run, size=9.5)


def _add_role_header(doc: Document, left: str, right: str = "") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(_clean(left))
    _set_run_font(run, size=10.0, bold=True)
    if right:
        run = p.add_run(f" | {_clean(right)}")
        _set_run_font(run, size=10.0, bold=True)


def _add_meta_line(doc: Document, location: str = "", dates: str = "") -> None:
    values = [value for value in (_clean(location), _clean(dates)) if value]
    if not values:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(" | ".join(values))
    _set_run_font(run, size=9.0)


def export_docx(master: dict, selection: dict, out_path: str | Path) -> Path:
    """Render an ATS-safe, text-only one-column DOCX."""
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(0)

    contact = master.get("contact", {})
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(_clean(contact.get("name", "")))
    _set_run_font(run, size=16, bold=True)

    headline = _clean(contact.get("headline", ""))
    if headline:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(headline)
        _set_run_font(run, size=10, bold=True)

    contact_parts = [
        _clean(contact.get("email", "")),
        _clean(contact.get("phone", "")),
        _clean(contact.get("location", "")),
    ]
    contact_parts.extend(_clean(link) for link in contact.get("links", []) if _clean(link))
    contact_parts = [part for part in contact_parts if part]
    if contact_parts:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(" | ".join(contact_parts))
        _set_run_font(run, size=8.8)

    summary = _clean(master.get("summary", ""))
    if summary:
        _add_section_heading(doc, "Summary")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(summary)
        _set_run_font(run, size=9.5)

    if selection.get("experience"):
        _add_section_heading(doc, "Experience")
        for selected in selection["experience"]:
            entry = selected["entry"]
            _add_role_header(doc, entry.get("company", ""), entry.get("title", ""))
            dates = " - ".join(
                value for value in (_clean(entry.get("start", "")), _clean(entry.get("end", ""))) if value
            )
            _add_meta_line(doc, entry.get("location", ""), dates)
            for bullet in selected["bullets"]:
                _add_bullet(doc, bullet["text"])

    if selection.get("projects"):
        _add_section_heading(doc, "Projects")
        for selected in selection["projects"]:
            entry = selected["entry"]
            _add_role_header(doc, entry.get("name", ""), entry.get("subtitle", ""))
            for bullet in selected["bullets"]:
                _add_bullet(doc, bullet["text"])

    education = master.get("education", [])
    if education:
        _add_section_heading(doc, "Education")
        for item in education:
            _add_role_header(doc, item.get("school", ""), item.get("degree", ""))
            dates = " - ".join(
                value for value in (_clean(item.get("start", "")), _clean(item.get("end", ""))) if value
            )
            _add_meta_line(doc, item.get("location", ""), dates)
            for detail in item.get("details", []):
                if _clean(detail):
                    _add_bullet(doc, detail)

    certifications = master.get("certifications", [])
    if certifications:
        _add_section_heading(doc, "Certifications")
        for cert in certifications:
            text = cert if isinstance(cert, str) else cert.get("name", "")
            if _clean(text):
                _add_bullet(doc, text)

    skills = selection.get("selected_verified_skills", []) or master.get("base_skills", [])
    if skills:
        _add_section_heading(doc, "Skills")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(", ".join(str(skill) for skill in skills))
        _set_run_font(run, size=9.5)

    # Avoid identifying/editor metadata. The file itself remains ordinary DOCX.
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = ""
    props.subject = ""
    props.keywords = ""
    props.comments = ""

    doc.save(out)
    return out


def audit_docx(path: str | Path) -> dict:
    """Check the generated package for hidden text/watermark-like constructs."""
    suspicious_markers = {
        "hidden_text": ("<w:vanish", "<w:webHidden"),
        "watermark_or_drawing": ("<w:pict", "<v:shape", "<w:background", "watermark"),
        "comments": ("comments.xml",),
    }
    found: dict[str, list[str]] = {key: [] for key in suspicious_markers}
    package = Path(path)
    with zipfile.ZipFile(package) as zf:
        names = zf.namelist()
        for name in names:
            lower_name = name.casefold()
            if "comments.xml" in lower_name:
                found["comments"].append(name)
            if not name.endswith((".xml", ".rels")):
                continue
            text = zf.read(name).decode("utf-8", errors="ignore")
            lower = text.casefold()
            for category, markers in suspicious_markers.items():
                for marker in markers:
                    if marker.casefold() in lower:
                        found[category].append(f"{name}:{marker}")
    clean = not any(found.values())
    return {
        "clean": clean,
        "findings": found,
        "checks": [
            "no hidden-text properties",
            "no watermark/background/VML drawing markers",
            "no comments part",
        ],
    }


def _serialize_selection(selection: dict) -> dict:
    """Remove duplicated full master entries from the JSON report."""
    report = {key: value for key, value in selection.items() if key not in ("experience", "projects")}
    report["experience"] = [
        {
            "entry_id": item["entry"].get("id"),
            "company": item["entry"].get("company"),
            "title": item["entry"].get("title"),
            "selected_bullet_ids": [bullet.get("id") for bullet in item["bullets"]],
            "bullet_scores": item["bullet_scores"],
        }
        for item in selection["experience"]
    ]
    report["projects"] = [
        {
            "entry_id": item["entry"].get("id"),
            "name": item["entry"].get("name"),
            "score": item["score"],
            "selected_bullet_ids": [bullet.get("id") for bullet in item["bullets"]],
            "bullet_scores": item["bullet_scores"],
        }
        for item in selection["projects"]
    ]
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a truthful ATS-safe DOCX for one JD")
    parser.add_argument("--master", required=True, help="Private locked master resume JSON")
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--jd", help="Job-description text file")
    source.add_argument("--jd-url", help="Public job-description URL")
    parser.add_argument("--out", required=True, help="Output .docx path")
    parser.add_argument("--report", help="Optional JSON selection/match report")
    parser.add_argument("--inventory", default=str(DEFAULT_INVENTORY), help="Skill inventory JSON")
    parser.add_argument("--max-skills", type=int, default=15)
    parser.add_argument("--max-exp-bullets", type=int, default=4)
    parser.add_argument("--max-projects", type=int, default=3)
    parser.add_argument("--max-project-bullets", type=int, default=3)
    args = parser.parse_args()

    master = load_master(args.master)
    inventory = load_inventory(args.inventory)
    jd_text = Path(args.jd).read_text(encoding="utf-8") if args.jd else fetch_jd_text(args.jd_url)

    selection = build_selection(
        master,
        jd_text,
        inventory,
        max_skills=args.max_skills,
        max_bullets_per_experience=args.max_exp_bullets,
        max_projects=args.max_projects,
        max_bullets_per_project=args.max_project_bullets,
    )
    out = export_docx(master, selection, args.out)
    audit = audit_docx(out)
    if not audit["clean"]:
        raise RuntimeError(f"generated DOCX failed safety audit: {audit['findings']}")

    report = _serialize_selection(selection)
    report["docx_audit"] = audit
    if args.report:
        report_path = Path(args.report)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")

    print(f"Created {out}")
    print(f"Selected {len(selection['selected_verified_skills'])} verified JD skills")
    print(
        "Selected experience bullets:",
        sum(len(item["bullets"]) for item in selection["experience"]),
    )
    print("Selected projects:", len(selection["projects"]))
    if selection["requested_but_unverified"]:
        print("JD gaps (not inserted):", ", ".join(selection["requested_but_unverified"]))
    print("DOCX hidden-text/watermark audit: clean")


if __name__ == "__main__":
    main()
